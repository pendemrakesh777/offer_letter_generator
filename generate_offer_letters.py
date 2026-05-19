from __future__ import annotations

import argparse
import csv
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = ROOT / "templates" / "offer_letter_template.docx"
DEFAULT_INPUT = ROOT / "data" / "offer_letters.csv"
DEFAULT_OUTPUT_DIR = ROOT / "generated_docs"

DEFAULT_COMPANY = "Adaptrix LLC"
DEFAULT_OFFICE = "6860 N Dallas Pkwy, Suite 200, Plano, TX 75024"


def format_offer_date(value: str) -> str:
    value = value.strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%B %d, %Y"):
        try:
            dt = datetime.strptime(value, fmt)
            return dt.strftime("%B %-d, %Y")
        except ValueError:
            continue
    return value


def format_salary(value: str) -> str:
    clean = value.strip().replace("$", "").replace(",", "")
    try:
        return f"${int(float(clean)):,}"
    except ValueError:
        return value.strip()


def safe_filename(name: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._ -]+", "", name).strip()
    return re.sub(r"\s+", "_", safe)


def clear_paragraph(paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    set_paragraph_text(new_para, text)
    if style is not None:
        new_para.style = style
    return new_para


def paragraph_texts(doc: Document) -> list[str]:
    return [paragraph.text.strip() for paragraph in doc.paragraphs]


def find_paragraph(doc: Document, startswith: str) -> int:
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip().startswith(startswith):
            return index
    raise ValueError(f"Could not find paragraph starting with: {startswith}")


def replace_responsibilities(doc: Document, responsibilities: list[str]) -> None:
    intro_index = find_paragraph(doc, "Below is a brief description")
    closing_index = find_paragraph(doc, "Work closely with internal teams")

    existing = doc.paragraphs[intro_index + 1 : closing_index]
    template_para = existing[0] if existing else doc.paragraphs[intro_index]

    for paragraph in existing:
        clear_paragraph(paragraph)

    for index, responsibility in enumerate(responsibilities):
        if index < len(existing):
            set_paragraph_text(existing[index], responsibility)
        else:
            template_para = insert_paragraph_after(
                template_para, responsibility, template_para.style
            )


def build_offer_letter(template: Path, output_path: Path, row: dict[str, str]) -> None:
    doc = Document(template)

    candidate_name = row["candidate_name"].strip()
    role = row["role"].strip()
    salary = format_salary(row["annual_salary"])
    offer_date = format_offer_date(row["offer_date"])
    company = row.get("company", "").strip() or DEFAULT_COMPANY
    office_address = row.get("office_address", "").strip() or DEFAULT_OFFICE
    signer_name = row.get("signer_name", "").strip() or "Mounika Bandari"
    signer_title = row.get("signer_title", "").strip() or "President"
    responsibilities = [
        item.strip()
        for item in row.get("responsibilities", "").split("|")
        if item.strip()
    ]

    date_index = find_paragraph(doc, "")
    for index, text in enumerate(paragraph_texts(doc)):
        if re.search(r"\b\d{4}\b", text):
            date_index = index
            break

    set_paragraph_text(doc.paragraphs[date_index], offer_date)
    set_paragraph_text(doc.paragraphs[find_paragraph(doc, "Dear ")], f"Dear {candidate_name},")

    main_index = find_paragraph(doc, "We are pleased to offer")
    main_text = (
        f"We are pleased to offer you a position with {company} for the role of {role}. "
        f"You will be employed by {company} at our corporate office, located at {office_address}. "
        f"Your annual salary will be {salary} per year. As part of your role, you may be "
        "required to work at client locations as assigned."
    )
    set_paragraph_text(doc.paragraphs[main_index], main_text)

    intro_index = find_paragraph(doc, "Below is a brief description")
    set_paragraph_text(
        doc.paragraphs[intro_index],
        f"Below is a brief description of the responsibilities involved in {role} position at our company",
    )

    if responsibilities:
        replace_responsibilities(doc, responsibilities)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "Mounika Bandari":
            set_paragraph_text(paragraph, signer_name)
        elif text == "President":
            set_paragraph_text(paragraph, signer_title)
        elif "Sravani Manne" in text and "Date" in text:
            set_paragraph_text(
                paragraph,
                f"____________________                                                                                                             ______________           {candidate_name}                                                                                                                                         Date",
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate offer letters from a CSV file.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    with args.input.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))

    if not rows:
        raise SystemExit(f"No rows found in {args.input}")

    for row in rows:
        candidate_name = row["candidate_name"].strip()
        output_path = args.output_dir / f"{safe_filename(candidate_name)}_offer_letter.docx"
        build_offer_letter(args.template, output_path, row)
        print(output_path)


if __name__ == "__main__":
    main()
